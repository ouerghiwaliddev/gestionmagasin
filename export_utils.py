"""Génération de rapports Excel et PDF sans état persistant."""

from io import BytesIO


def _text(value):
    return "" if value is None else str(value)


def build_excel(sections):
    """Construit un classeur XLSX, avec une feuille par section."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles = set()

    for section in sections:
        base_title = section["title"][:31] or "Données"
        title = base_title
        suffix = 2
        while title in used_titles:
            title = f"{base_title[:27]} {suffix}"
            suffix += 1
        used_titles.add(title)

        sheet = workbook.create_sheet(title)
        sheet.append(section["headers"])
        for row in section["rows"]:
            sheet.append([_text(value) for value in row])

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
            cell.alignment = Alignment(horizontal="center")

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            max_length = max(len(_text(cell.value)) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 12), 45)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _pdf_escape(value):
    encoded = _text(value).encode("cp1252", errors="replace").decode("latin-1")
    return encoded.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _column_widths(headers, rows, available_width):
    weights = []
    for index, header in enumerate(headers):
        values = [_text(row[index]) for row in rows[:100] if index < len(row)]
        weights.append(min(max([len(_text(header)), *(len(value) for value in values)], default=8), 28))
    total = sum(weights) or 1
    return [available_width * weight / total for weight in weights]


def _table_page(title, headers, rows, page_number, continued=False):
    width, height = 842, 595
    margin, table_width = 30, 782
    top, row_height = 540, 15
    font_size = 7 if len(headers) <= 8 else 6
    column_widths = _column_widths(headers, rows, table_width)
    commands = [
        "0.12 0.20 0.35 rg",
        f"BT /F2 16 Tf {margin} 565 Td ({_pdf_escape(title)}) Tj ET",
        "0.35 0.40 0.48 rg",
        f"BT /F1 8 Tf {margin} 550 Td (Page {page_number}{' - suite' if continued else ''}) Tj ET",
    ]

    def draw_row(values, y, header=False):
        x = margin
        if header:
            commands.append(f"0.15 0.39 0.82 rg {margin} {y} {table_width} {row_height} re f")
            commands.append("1 1 1 rg")
            font = "F2"
        else:
            commands.append("0.15 0.18 0.22 rg")
            font = "F1"

        for index, cell in enumerate(values):
            cell_width = column_widths[index]
            max_chars = max(3, int((cell_width - 6) / (font_size * 0.52)))
            displayed = _text(cell)
            if len(displayed) > max_chars:
                displayed = displayed[:max(1, max_chars - 3)] + "..."
            commands.append(
                f"BT /{font} {font_size} Tf {x + 3:.2f} {y + 4:.2f} Td ({_pdf_escape(displayed)}) Tj ET"
            )
            commands.append(f"0.72 0.76 0.82 RG {x:.2f} {y:.2f} {cell_width:.2f} {row_height} re S")
            x += cell_width

    draw_row(headers, top, header=True)
    y = top - row_height
    if rows:
        for row in rows:
            draw_row(row, y)
            y -= row_height
    else:
        commands.append(f"0.35 0.40 0.48 rg BT /F1 10 Tf {margin} {y} Td (Aucune donnée) Tj ET")

    return "\n".join(commands).encode("latin-1", errors="replace")


def build_pdf(sections):
    """Construit un PDF paysage paginé avec une table par section."""
    streams = []
    page_number = 1
    rows_per_page = 32

    for section in sections:
        rows = list(section["rows"])
        chunks = [rows[index:index + rows_per_page] for index in range(0, len(rows), rows_per_page)] or [[]]
        for index, chunk in enumerate(chunks):
            streams.append(_table_page(
                section["title"], section["headers"], chunk,
                page_number, continued=index > 0
            ))
            page_number += 1

    page_ids = [5 + index * 2 for index in range(len(streams))]
    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Count {len(page_ids)} /Kids [{' '.join(f'{page_id} 0 R' for page_id in page_ids)}] >>".encode(),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        4: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>",
    }
    for index, stream in enumerate(streams):
        page_id = page_ids[index]
        content_id = page_id + 1
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 842 595] "
            f"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        objects[content_id] = f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id in range(1, max(objects) + 1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode())
        output.extend(objects[object_id])
        output.extend(b"\nendobj\n")

    xref_offset = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode()
    )
    return BytesIO(output)
